from __future__ import annotations

from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfgen import canvas

from offline_docs import write_offline_documents

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "backend" / "data" / "sample_docs"

MD = """# WebSocket 连接指南

## 连接生命周期

WebSocket 连接默认空闲超时为 10 分钟。客户端应每 30 秒发送一次 ping，服务端在收到 ping 后返回 pong。

## 断线重连

连接中断后，客户端应使用指数退避策略重连，初始等待 1 秒，最大等待 30 秒。重连请求必须携带原请求的 request_id，服务端据此恢复订阅关系。

```python
await websocket.send_json({"type": "ping"})
```

## 常见错误

- 错误码 4010：鉴权过期，需要重新获取 access_token。
- 错误码 4011：请求频率超限，应在 Retry-After 后再连接。
- 错误码 4012：订阅主题不存在，需要检查 topic 参数。
"""

ERRORS = """# API 错误码手册

## 4010 鉴权过期

access_token 有效期 2 小时。刷新 token 后应使用新的 Authorization 头重新发起请求。

## 4290 频率限制

默认每个应用每分钟 600 次请求。超限时响应头会返回 Retry-After，建议使用指数退避。

## 5003 上游超时

服务端依赖的上游接口超时。客户端可以重试一次；第二次仍失败时请提交工单并提供 request_id。
"""


def create_markdown() -> None:
    (OUT / "websocket-guide.md").write_text(MD, encoding="utf-8")
    (OUT / "api-errors.md").write_text(ERRORS, encoding="utf-8")


def create_pdf() -> None:
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    path = OUT / "architecture.pdf"
    c = canvas.Canvas(str(path))
    c.setFont("STSong-Light", 16)
    lines = [
        "DocFlow 技术架构说明",
        "文档解析：PDF、PPTX、Markdown、图片",
        "向量存储：Milvus 稠密向量检索",
        "混合检索：BM25 与向量分数融合",
        "大模型：通义千问 qwen-max 流式输出",
    ]
    y = 810
    for line in lines:
        c.drawString(60, y, line)
        y -= 34
    c.save()


def create_pptx() -> None:
    prs = Presentation()
    slides = [
        ("版本发布说明 v2.4", "2026 年 9 月发布\n新增多知识库检索、流式回答和文档引用"),
        ("兼容性变更", "旧版 /v1/chat 接口将在 2027-01-01 下线\n请迁移到 /v2/chat 流式接口"),
        ("故障排查", "如果检索结果为空：检查文档状态、知识库选择和向量索引"),
    ]
    for title, body in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
    prs.save(OUT / "release-notes.pptx")


def create_image() -> None:
    image = Image.new("RGB", (1000, 620), "#f8faf9")
    draw = ImageDraw.Draw(image)
    font_path = Path("C:/Windows/Fonts/msyh.ttc")
    font = ImageFont.truetype(str(font_path), 34) if font_path.exists() else ImageFont.load_default()
    small = ImageFont.truetype(str(font_path), 24) if font_path.exists() else ImageFont.load_default()
    draw.rounded_rectangle((55, 55, 945, 565), radius=28, fill="#ffffff", outline="#cbd9d3", width=2)
    draw.text((95, 110), "错误排查流程", fill="#0f766e", font=font)
    steps = [
        "1. 检查 access_token 是否过期（错误码 4010）",
        "2. 检查请求频率是否超过 600 次/分钟",
        "3. 检查 WebSocket 是否每 30 秒发送 ping",
        "4. 断线后使用指数退避并携带 request_id",
    ]
    y = 200
    for step in steps:
        draw.text((100, y), step, fill="#26352f", font=small)
        y += 72
    image.save(OUT / "troubleshooting.png")


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    write_offline_documents(OUT)
    create_markdown()
    create_pdf()
    create_pptx()
    create_image()
    for path in sorted(OUT.iterdir()):
        print(path.name, path.stat().st_size)


if __name__ == "__main__":
    main()
