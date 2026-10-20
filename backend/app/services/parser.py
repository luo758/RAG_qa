from __future__ import annotations

from pathlib import Path

from pypdf import PdfReader
from pptx import Presentation

from .vision import get_vision_provider

TEXT_EXTENSIONS = {".md", ".markdown", ".txt"}
PDF_EXTENSIONS = {".pdf"}
PPT_EXTENSIONS = {".pptx"}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".bmp"}
SUPPORTED_EXTENSIONS = TEXT_EXTENSIONS | PDF_EXTENSIONS | PPT_EXTENSIONS | IMAGE_EXTENSIONS


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in TEXT_EXTENSIONS:
        return path.read_text(encoding="utf-8", errors="replace")
    if suffix in PDF_EXTENSIONS:
        reader = PdfReader(str(path))
        return "\n\n".join(page.extract_text() or "" for page in reader.pages)
    if suffix in PPT_EXTENSIONS:
        presentation = Presentation(str(path))
        parts: list[str] = []
        for slide_index, slide in enumerate(presentation.slides, start=1):
            parts.append(f"# 第 {slide_index} 页")
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    parts.append(shape.text.strip())
            try:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    parts.append(f"备注：{notes}")
            except Exception:  # noqa: BLE001
                pass
        return "\n\n".join(parts)
    if suffix in IMAGE_EXTENSIONS:
        return get_vision_provider().describe_image(path)
    raise ValueError(f"不支持的文件类型：{suffix}")
